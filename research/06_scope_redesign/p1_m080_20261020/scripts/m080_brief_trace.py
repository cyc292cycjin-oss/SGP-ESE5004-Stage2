"""M08-0 item 6 check: every data cell of DATA_INVENTORY_BRIEF.xlsx appears verbatim in DATA_INVENTORY_v2.xlsx (any sheet).
Rows with a single non-empty cell (titles / captions) are skipped. Output <out>/evidence/M080_BRIEF_TRACE.json.

    python -I m080_brief_trace.py --meet /mnt/d/.../p1_meet_20261020 --out /mnt/d/.../p1_m080_20261020
"""
from __future__ import annotations

import argparse
import hashlib
import json
from pathlib import Path

import openpyxl


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--meet", type=Path, required=True)
    ap.add_argument("--out", type=Path, required=True)
    ap.add_argument("--detail", default="DATA_INVENTORY_v2.xlsx", help="detail workbook (supplement 01: DATA_INVENTORY_v3.xlsx)")
    ap.add_argument("--name", default="M080_BRIEF_TRACE")
    a = ap.parse_args()
    bpath, vpath = a.meet / "DATA_INVENTORY_BRIEF.xlsx", a.meet / a.detail
    b, v = openpyxl.load_workbook(bpath), openpyxl.load_workbook(vpath)
    vtexts = {str(c) for ws in v.worksheets for row in ws.iter_rows(values_only=True) for c in row if c is not None}
    checked, missing, per_sheet = 0, [], {}
    for ws in b.worksheets:
        n = 0
        for i, row in enumerate(ws.iter_rows(values_only=True), 1):
            cells = [c for c in row if c is not None]
            if len(cells) <= 1:
                continue
            if ws.title != "说明与统计" and i == 1:
                continue   # header row of the brief tables (column names)
            for c in cells:
                n += 1
                if str(c) not in vtexts:
                    missing.append(dict(sheet=ws.title, row=i, text=str(c)[:120]))
        per_sheet[ws.title] = n
        checked += n
    out = dict(brief=str(bpath), brief_sha256=hashlib.sha256(bpath.read_bytes()).hexdigest(), detail=str(vpath),
               detail_sha256=hashlib.sha256(vpath.read_bytes()).hexdigest(), cells_checked=checked, per_sheet=per_sheet,
               not_found_verbatim=len(missing), missing=missing[:50])
    a.out.joinpath("evidence").mkdir(parents=True, exist_ok=True)
    (a.out / "evidence" / f"{a.name}.json").write_text(json.dumps(out, indent=1, ensure_ascii=False) + "\n", encoding="utf-8")
    print(json.dumps({k: v for k, v in out.items() if k != "missing"}, indent=1, ensure_ascii=False))
    for m in missing[:20]:
        print("MISSING", m)


if __name__ == "__main__":
    main()
