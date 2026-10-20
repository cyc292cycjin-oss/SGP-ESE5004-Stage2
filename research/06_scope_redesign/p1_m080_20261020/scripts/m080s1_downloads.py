"""M08-0 supplement 01 section 3 (item 131), read-only: the four user downloads for M08 and the technology-data v0.15.0 HICP check.

Reads work/p1_s4r_downloads_20261020/{R1_demand,R6_costs}/... and the technology-data v0.15.0 source tarball (registry TARBALL_SHA
c4787c1d...); writes <out>/evidence/M080S1_DOWNLOADS.json. Nothing is changed, built or solved.

    python -I m080s1_downloads.py --out /mnt/d/ResearchWorkspaces/ASEAN/work/p1_m080_20261020
"""
from __future__ import annotations

import argparse
import hashlib
import io
import json
import tarfile
import warnings
from pathlib import Path

import openpyxl
import pandas as pd

warnings.filterwarnings("ignore")
DL = Path("/mnt/d/ResearchWorkspaces/ASEAN/work/p1_s4r_downloads_20261020")
FILES = {"cement": DL / "R1_demand/myb1-2023-cemen-ERT.xlsx", "steel": DL / "R1_demand/exp-2026-10-20_15_02_22.xlsx",
         "population": DL / "R1_demand/unpopulation_dataportal_20261020224306.csv", "hicp": DL / "R6_costs/prc_hicp_aind$defaultview_spreadsheet.xlsx"}
NOT_USED = {"R6_costs/ESTAT_PRC_HICP_AIND$DEFAULTVIEW_1.0.xml.gz": "data structure only (no values), not used",
            "R1_demand/exp-2026-10-20_14_37_36.png": "earlier screenshot, not in the supplement list, not registered"}
TD_TAR = Path("/home/jin/s1_tmp/dl/td015/v0.15.0.tar.gz")
TD_TAR_SHA = "c4787c1d72b9ec9fae032d37f906a4d6c92d4adfe111c55680f59f62bb162ccb"
USGS = {"Brunei": "BN", "Burma4": "MM", "Cambodia": "KH", "Indonesia": "ID", "Laos": "LA", "Malaysia": "MY", "Philippines": "PH",
        "Singapore": "SG", "Thailand": "TH", "Timor-Leste": "TL", "Vietnam": "VN"}
WSA = {"Brunei Darussalam": "BN", "Brunei": "BN", "Cambodia": "KH", "Indonesia": "ID", "Lao PDR": "LA", "Laos": "LA", "Malaysia": "MY", "Myanmar": "MM",
       "Philippines": "PH", "Singapore": "SG", "Thailand": "TH", "Timor-Leste": "TL", "Viet Nam": "VN"}
ASEAN = ["BN", "ID", "KH", "LA", "MM", "MY", "PH", "SG", "TH", "TL", "VN"]
HICP_ROWS = {"EU27": "European Union - 27 countries (from 2020)", "EA20": "Euro area – 20 countries (2023-2025)",
             "EEA": "European Economic Area (EEA18-1995, EEA28-2004, EEA30-2007, EEA31-2013, EEA30-2020)"}


def sha(p: Path | bytes) -> str:
    return hashlib.sha256(p if isinstance(p, bytes) else p.read_bytes()).hexdigest()


def cement():
    wb = openpyxl.load_workbook(FILES["cement"], read_only=True, data_only=True)
    note = " ".join(str(c) for row in wb["Note"].iter_rows(values_only=True) for c in row if c)
    rows = list(wb["T22"].iter_rows(values_only=True))
    title = " ".join(str(c) for r in rows[:4] for c in r if c)
    hdr = next(r for r in rows if r and r[0] == "Country or locality")
    years = [int(c) for c in hdr if isinstance(c, (int, float)) or (isinstance(c, str) and c.isdigit())]
    vals = {}
    for r in rows:
        if r and r[0] in USGS:
            nums = [r[i] for i in range(2, len(r), 2)][: len(years)]
            flags = [r[i] for i in range(3, len(r), 2)][: len(years)]
            vals[USGS[r[0]]] = {str(y): dict(kt=v, flag=(f or "")) for y, v, f in zip(years, nums, flags)}
    foot = [" ".join(str(c) for c in r if c) for r in rows if r and isinstance(r[0], str) and r[0][:1].isdigit()]
    return dict(title=title, note=note[:400], years=years, unit="thousand metric tons", asean=vals, missing=[c for c in ASEAN if c not in vals],
                footnotes=foot[:4], properties=dict(title=wb.properties.title, creator=wb.properties.creator, modified=str(wb.properties.modified)))


def steel():
    ws = openpyxl.load_workbook(FILES["steel"], read_only=True, data_only=True).active
    rows = list(ws.iter_rows(values_only=True))
    hdr = next(r for r in rows if r and r[0] == "国家")
    years = [int(c) for c in hdr[1:] if c]
    vals = {WSA[r[0]]: {str(y): v for y, v in zip(years, r[1:])} for r in rows if r and r[0] in WSA}
    return dict(sheet=ws.title, title=rows[0][0], years=years, unit_in_file=None,
                unit_inferred="thousand tonnes (VN 2024 21,974 = worldsteel WSIF 2025 VN 22.0 Mt quoted in S4R)", asean=vals,
                missing=[c for c in ASEAN if c not in vals])


def population():
    df = pd.read_csv(FILES["population"])
    r = df[(df.Iso2 == "TL") & (df.Time == 2050)].iloc[0]
    return dict(indicator=r.IndicatorName, source=f"{r.Source} {r.SourceYear}", author=r.Author, location=r.Location, variant=r.Variant, sex=r.Sex,
                estimate_method=r.EstimateMethod, year=int(r.Time), value=float(r.Value), years_in_file=f"{int(df.Time.min())}-{int(df.Time.max())}")


def hicp_sheet(src, sheet, header_row_label="TIME"):
    ws = openpyxl.load_workbook(src, data_only=True)[sheet]
    rows = list(ws.iter_rows(values_only=True))
    meta = [" | ".join(str(c) for c in r if c is not None) for r in rows[:7]]
    hdr = next(r for r in rows if r and r[0] == header_row_label)
    years = [str(c) for c in hdr[1:] if c]
    out = {}
    for k, label in HICP_ROWS.items():
        r = next((r for r in rows if r and r[0] == label), None)
        if r:
            out[k] = {y: (None if v in (":", None) else float(v)) for y, v in zip(years, r[1:][::2])}
    return meta, out


def hicp():
    meta1, idx = hicp_sheet(FILES["hicp"], "Sheet 1")
    meta6, rate = hicp_sheet(FILES["hicp"], "Sheet 6")
    ratio = {k: round(v["2025"] / v["2020"], 5) for k, v in idx.items()}
    cum = {}
    for k, v in rate.items():
        f = 1.0
        for y in range(2021, 2026):
            f *= 1 + v[str(y)] / 100.0
        cum[k] = round(f, 5)
    # technology-data v0.15.0: code + its own inflation input file
    assert sha(TD_TAR) == TD_TAR_SHA
    with tarfile.open(TD_TAR) as t:
        helpers = t.extractfile("technology-data-0.15.0/scripts/_helpers.py").read()
        cfg = t.extractfile("technology-data-0.15.0/config.yaml").read().decode()
        inp_name = "technology-data-0.15.0/inputs/prc_hicp_aind__custom_20097956_spreadsheet.xlsx"
        inp = t.extractfile(inp_name).read()
        outputs = sorted(n for n in t.getnames() if n.startswith("technology-data-0.15.0/outputs/costs_") and n.endswith(".csv"))
    lines = helpers.decode().splitlines()
    row_line = next(i for i, l in enumerate(lines, 1) if "row_to_use = \"European" in l)
    df = pd.read_excel(io.BytesIO(inp), sheet_name="Sheet 1", index_col=0, na_values=[":", "d"], header=8)
    td_rates = {}
    for k, label in HICP_ROWS.items():
        if label in df.index:
            s = df.loc[label].dropna()
            td_rates[k] = {str(int(c)): float(s[c]) for c in s.index if str(c).isdigit() and 2021 <= int(c) <= 2025}
    td_factor = {}
    for k, v in td_rates.items():
        f = 1.0
        for y in range(2021, 2026):
            f *= 1 + v[str(y)] / 100.0
        td_factor[k] = round(f, 5)
    return dict(download=dict(sheet1=meta1, sheet6=meta6, index=idx, index_ratio_2025_2020=ratio, rates_2021_2025={k: {y: v[y] for y in map(str, range(2021, 2026))} for k, v in rate.items()},
                              rate_cumprod_2021_2025=cum),
                technology_data_v0_15_0=dict(tarball_sha256=TD_TAR_SHA, helpers_sha256=sha(helpers), row_to_use_line=row_line, row_to_use=lines[row_line - 1].strip(),
                                             method="prepare_inflation_rate: Sheet 1 of the input file, annual average rate of change / 100; adjust_for_inflation: "
                                                    "factor = prod(1 + r_y) for y = ref_year+1 .. eur_year",
                                             eur_year_config=[l.strip() for l in cfg.splitlines() if l.strip().startswith("eur_year")],
                                             output_years_config=[l.strip() for l in cfg.splitlines() if l.strip().startswith("years:")],
                                             outputs_in_tarball=[n.split("/")[-1] for n in outputs], input_file=inp_name.split("/")[-1], input_sha256=sha(inp),
                                             input_rates_2021_2025=td_rates, factor_2020_2025=td_factor))


def td_by_year():
    """Supplement 01 section 4 (record only): conventional-plant parameters of technology-data v0.15.0 by output year (tarball outputs/costs_<year>.csv)."""
    techs = ["CCGT", "OCGT", "coal", "lignite", "oil", "nuclear", "biomass"]
    out = {}
    with tarfile.open(TD_TAR) as t:
        for y in (2020, 2025, 2030, 2035, 2040, 2045, 2050):
            df = pd.read_csv(io.BytesIO(t.extractfile(f"technology-data-0.15.0/outputs/costs_{y}.csv").read()))
            sub = df[df.technology.isin(techs) & df.parameter.isin(["efficiency", "FOM", "lifetime"])]
            out[str(y)] = sub.pivot_table(index="technology", columns="parameter", values="value", aggfunc="first").round(4).to_dict(orient="index")
    return out


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--out", type=Path, required=True)
    a = ap.parse_args()
    out = dict(schema="m080s1-downloads-v1", solver_runs=0,
               files={k: dict(path=str(p), name=p.name, bytes=p.stat().st_size, sha256=sha(p)) for k, p in FILES.items()},
               not_registered={k: dict(sha256=sha(DL / k), note=v) for k, v in NOT_USED.items() if (DL / k).exists()},
               not_registered_elsewhere="mis-202606-cemen.xlsx (USGS Mineral Industry Surveys, US monthly) is in the user's Downloads folder, not in work/: wrong / not useful, not used",
               cement=cement(), steel=steel(), population=population(), hicp=hicp(), technology_data_v0_15_0_by_year=td_by_year())
    (a.out / "evidence").mkdir(parents=True, exist_ok=True)
    (a.out / "evidence/M080S1_DOWNLOADS.json").write_text(json.dumps(out, indent=1, ensure_ascii=False, default=str) + "\n", encoding="utf-8")
    h = out["hicp"]
    print(json.dumps(dict(files={k: v["sha256"][:16] for k, v in out["files"].items()}, cement_missing=out["cement"]["missing"], steel_missing=out["steel"]["missing"],
                          tl_2050=out["population"]["value"], index_ratio=h["download"]["index_ratio_2025_2020"], rate_cumprod=h["download"]["rate_cumprod_2021_2025"],
                          td=dict(row=h["technology_data_v0_15_0"]["row_to_use"], factor=h["technology_data_v0_15_0"]["factor_2020_2025"],
                                  outputs=h["technology_data_v0_15_0"]["outputs_in_tarball"])), indent=1, ensure_ascii=False))


if __name__ == "__main__":
    main()
