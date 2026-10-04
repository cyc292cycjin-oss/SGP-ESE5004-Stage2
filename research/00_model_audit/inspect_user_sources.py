"""Bounded read of supplied source metadata; no workbook/PDF edits."""
from pathlib import Path
import json,hashlib
from openpyxl import load_workbook
from pypdf import PdfReader
out=Path(__file__).resolve().parent
base=Path('C:/Users/20122/Desktop/Singapore')
records=[]
for p in sorted((base/'技术目录').glob('*.xlsx')):
    w=load_workbook(p,read_only=True,data_only=True)
    rec={'file':str(p),'sha256':hashlib.sha256(p.read_bytes()).hexdigest(),'sheets':w.sheetnames,'properties':{'title':w.properties.title,'created':str(w.properties.created),'modified':str(w.properties.modified)},'cells':{}}
    selected=[s for s in w.sheetnames if '86' in s and ('AEC' in s or 'Electro' in s)]
    for s in selected:
        cells=[]
        for row in w[s].iter_rows(min_row=1,max_row=90,max_col=18):
            values={c.coordinate:c.value for c in row if c.value is not None}
            if values:cells.append(values)
        rec['cells'][s]=cells
    records.append(rec);w.close()
papers=[]
for p in sorted((base/'文献Stage2').glob('*.pdf')):
    r=PdfReader(p)
    papers.append({'file':str(p),'pages':len(r.pages),'metadata':str(r.metadata),'first_page_text':r.pages[0].extract_text()})
(out/'USER_SOURCE_INVENTORY.json').write_text(json.dumps({'workbooks':records,'paper_candidates':papers},ensure_ascii=False,indent=2,default=str)+'\n',encoding='utf-8')
print(json.dumps({'workbooks':[{'file':r['file'],'selected_sheets':list(r['cells'])} for r in records],'papers':[{'file':r['file'],'first_page':r['first_page_text'][:700]} for r in papers]},ensure_ascii=True))
